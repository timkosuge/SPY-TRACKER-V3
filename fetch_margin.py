"""FINRA margin statistics: parse FINRA's workbook, fetch it, and keep every month in margin_monthly.

FINRA publishes the monthly figures only on its page and as an Excel download; it offers no feed. The pipeline asks for
the download once a day, directly and then through the site's route on Cloudflare. Every attempt is logged with its result, so the site can say when a month could not be fetched
instead of showing an old figure as the latest.
"""
import csv
import io
import sqlite3
import sys
from datetime import datetime

import pytz
import requests

DB_PATH = "spy_data.db"
HISTORY_CSV = "data/finra_margin_history.csv"
FINRA_XLSX = "https://www.finra.org/sites/default/files/2021-03/margin-statistics.xlsx"
FINRA_PAGE = "https://www.finra.org/rules-guidance/key-topics/margin-accounts/margin-statistics"
SITE_ROUTE = "https://spy-tracker-v3.pages.dev/finra-margin"
HEADER = "Year-Month"
ET = pytz.timezone("America/New_York")


def parse_workbook(content):
    """Every month in FINRA's workbook as (month, debit, free_cash, free_margin), millions of dollars, oldest first."""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    if not rows or str(rows[0][0]).strip() != HEADER:
        raise ValueError(f"unexpected first cell {rows[0][0]!r}" if rows else "empty workbook")
    out = []
    for r in rows[1:]:
        month = str(r[0]).strip() if r and r[0] is not None else ""
        if len(month) != 7 or month[4] != "-" or r[1] is None:
            continue
        num = lambda v: int(v) if v not in (None, "") else None
        out.append((month, num(r[1]), num(r[2]), num(r[3])))
    out.sort()
    return out


def read_history_csv(path=HISTORY_CSV):
    with open(path, newline="") as f:
        return [(r["month"], int(r["debit"]), int(r["free_cash"]) if r["free_cash"] else None, int(r["free_margin"]) if r["free_margin"] else None)
                for r in csv.DictReader(f)]


def write_history_csv(rows, path=HISTORY_CSV):
    with open(path, "w", newline="") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(["month", "debit", "free_cash", "free_margin"])
        for m, d, fc, fm in rows:
            w.writerow([m, d, "" if fc is None else fc, "" if fm is None else fm])


def create(conn):
    conn.execute("CREATE TABLE IF NOT EXISTS margin_monthly (month TEXT PRIMARY KEY, debit INTEGER, free_cash INTEGER, free_margin INTEGER, source TEXT)")
    conn.execute("CREATE TABLE IF NOT EXISTS margin_fetch_log (attempted_at TEXT PRIMARY KEY, http_status INTEGER, ok INTEGER, months INTEGER, latest TEXT, detail TEXT)")


def _workbook_from(url, headers):
    """One request. Returns (rows or None, http status, what happened)."""
    try:
        r = requests.get(url, timeout=45, headers=headers)
    except requests.RequestException as e:
        return None, 0, f"request failed: {e}"
    if r.status_code == 200 and r.content[:2] == b"PK":
        try:
            return parse_workbook(r.content), 200, "workbook fetched"
        except Exception as e:
            return None, 200, f"workbook unreadable: {e}"
    if r.status_code == 502 and r.content[:1] == b"{":
        try:
            import json
            body = json.loads(r.content)
            return None, body.get("finra_status") or 502, body.get("error") or "the site's route could not reach FINRA"
        except ValueError:
            pass
    if r.status_code != 200:
        return None, r.status_code, f"FINRA answered HTTP {r.status_code}"
    return None, r.status_code, "FINRA answered with a page, not the workbook"


def fetch():
    """FINRA's workbook, asked for directly and then through the site's own route on Cloudflare. Returns (rows, status, detail)."""
    direct = _workbook_from(FINRA_XLSX, {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0 Safari/537.36",
        "Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,*/*",
        "Referer": FINRA_PAGE})
    if direct[0]:
        return direct[0], direct[1], "workbook fetched from FINRA directly"
    routed = _workbook_from(SITE_ROUTE, {"User-Agent": "spy-tracker-pipeline"})
    if routed[0]:
        return routed[0], routed[1], "workbook fetched through the site's Cloudflare route"
    return None, routed[1] or direct[1], f"refused on both networks: GitHub got '{direct[2]}', Cloudflare got '{routed[2]}'"


def main():
    try:
        with open(DB_PATH, "rb") as f:
            if f.read(16) != b"SQLite format 3\x00":
                print("margin: no database here; nothing stored")
                return
    except OSError:
        print("margin: no database here; nothing stored")
        return
    rows, status, detail = fetch()
    conn = sqlite3.connect(DB_PATH)
    create(conn)
    if rows:
        conn.executemany("INSERT OR REPLACE INTO margin_monthly VALUES (?,?,?,?,'finra_xlsx')", rows)
    now = datetime.now(ET).isoformat(timespec="seconds")
    conn.execute("INSERT OR REPLACE INTO margin_fetch_log VALUES (?,?,?,?,?,?)",
                 (now, status, 1 if rows else 0, len(rows or []), rows[-1][0] if rows else None, detail))
    conn.commit()
    conn.close()
    print(f"margin: {detail}" + (f", {len(rows)} months through {rows[-1][0]}" if rows else ""))


if __name__ == "__main__":
    if len(sys.argv) == 3 and sys.argv[1] == "--seed-csv":
        with open(sys.argv[2], "rb") as f:
            rows = parse_workbook(f.read())
        write_history_csv(rows)
        print(f"{HISTORY_CSV}: {len(rows)} months, {rows[0][0]} → {rows[-1][0]}")
    else:
        main()
